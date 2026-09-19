import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = Presentation()
    # 16:9 Widescreen layout
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette: Modern Cyber Dark
    C_BG = RGBColor(10, 14, 23)           # #0A0E17 Deep Midnight Navy
    C_CARD = RGBColor(18, 24, 38)         # #121826 Slate Card Surface
    C_BORDER_BLUE = RGBColor(14, 165, 233)# #0EA5E9 Sky Blue
    C_CYAN = RGBColor(6, 182, 212)        # #06B6D4 Electric Cyan
    C_GREEN = RGBColor(16, 185, 129)      # #10B981 Emerald Green
    C_AMBER = RGBColor(245, 158, 11)      # #F59E0B Amber
    C_ROSE = RGBColor(244, 63, 94)        # #F43F5E Coral Rose
    C_PURPLE = RGBColor(139, 92, 246)     # #8B5CF6 Violet
    C_WHITE = RGBColor(248, 250, 252)     # #F8FAFC Pure Light
    C_TEXT_SECONDARY = RGBColor(148, 163, 184) # #94A3B8 Slate Text
    C_TEXT_MUTED = RGBColor(100, 116, 139)     # #64748B Muted Text

    def set_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = C_BG
        bg.line.fill.background()
        return bg

    def add_header(slide, section_tag, title_text, subtitle_text=None):
        # Top Accent Badge
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(10), Inches(0.35))
        tf_tag = tag_box.text_frame
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = section_tag.upper()
        p_tag.font.size = Pt(11)
        p_tag.font.bold = True
        p_tag.font.color.rgb = C_CYAN

        # Slide Main Title
        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.5), Inches(0.75))
        tf_t = t_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(25)
        p_t.font.bold = True
        p_t.font.color.rgb = C_WHITE

        # Subtitle
        if subtitle_text:
            s_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.35), Inches(11.5), Inches(0.45))
            tf_s = s_box.text_frame
            tf_s.word_wrap = True
            p_s = tf_s.paragraphs[0]
            p_s.text = subtitle_text
            p_s.font.size = Pt(12)
            p_s.font.color.rgb = C_TEXT_SECONDARY

    def add_card(slide, left, top, width, height, title, items, border_color=C_BORDER_BLUE, title_color=C_CYAN):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)

        tx = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), width - Inches(0.5), height - Inches(0.4))
        tf = tx.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.size = Pt(15)
        p0.font.bold = True
        p0.font.color.rgb = title_color
        p0.space_after = Pt(10)

        for it in items:
            p = tf.add_paragraph()
            p.text = f"•  {it}"
            p.font.size = Pt(11)
            p.font.color.rgb = C_WHITE
            p.space_after = Pt(6)

    # =========================================================================
    # SLIDE 1: PROJECT NAME & COVER
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_bg(s1)

    # Accent decorative bar
    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.5), Inches(1.5), Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = C_CYAN
    bar.line.fill.background()

    # Track Tag
    tr_box = s1.shapes.add_textbox(Inches(0.8), Inches(1.75), Inches(11.5), Inches(0.5))
    p_tr = tr_box.text_frame.paragraphs[0]
    p_tr.text = "AGENTIC AI TRACK  |  HACKATHON 2026 SUBMISSION"
    p_tr.font.size = Pt(12)
    p_tr.font.bold = True
    p_tr.font.color.rgb = C_CYAN

    # Project Name
    pn_box = s1.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.7), Inches(2.0))
    tf_pn = pn_box.text_frame
    tf_pn.word_wrap = True
    p_pn = tf_pn.paragraphs[0]
    p_pn.text = "DevOps Autopilot"
    p_pn.font.size = Pt(46)
    p_pn.font.bold = True
    p_pn.font.color.rgb = C_WHITE

    p_pn2 = tf_pn.add_paragraph()
    p_pn2.text = "Autonomous Infra Waste Resolver"
    p_pn2.font.size = Pt(30)
    p_pn2.font.bold = True
    p_pn2.font.color.rgb = C_BORDER_BLUE

    # Pitch Subtitle
    sub_box = s1.shapes.add_textbox(Inches(0.8), Inches(4.5), Inches(11.5), Inches(1.1))
    tf_sub = sub_box.text_frame
    tf_sub.word_wrap = True
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = "An intelligent multi-agent AI system that detects cloud waste across Kubernetes & AWS, simulates capacity risk, enforces strict safety policies, and executes risk-governed remediations with interactive Slack ChatOps."
    p_sub.font.size = Pt(13)
    p_sub.font.color.rgb = C_TEXT_SECONDARY

    # Metadata Badges Box
    meta_box = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.8), Inches(11.7), Inches(0.9))
    meta_box.fill.solid()
    meta_box.fill.fore_color.rgb = C_CARD
    meta_box.line.color.rgb = C_BORDER_BLUE
    meta_box.line.width = Pt(1)

    tx_m = s1.shapes.add_textbox(Inches(1.0), Inches(5.95), Inches(11.3), Inches(0.6))
    p_m = tx_m.text_frame.paragraphs[0]
    p_m.text = "⚡ Core Stack: Python 3.10+, Flask API, SQLAlchemy, Three.js WebGL, Slack Block Kit, OpenAI Function Calling"
    p_m.font.size = Pt(11)
    p_m.font.bold = True
    p_m.font.color.rgb = C_WHITE

    # =========================================================================
    # SLIDE 2: PROJECT DESCRIPTION
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_bg(s2)
    add_header(s2, "Project Description", "What is DevOps Autopilot?", "An autonomous multi-agent copilot engineered to eliminate cloud overprovisioning safely.")

    add_card(s2, Inches(0.8), Inches(2.1), Inches(3.6), Inches(4.7), "Autonomous SRE Copilot", [
        "Continuous 24/7 background telemetry ingestion across K8s and AWS compute.",
        "Identifies idle pods, over-allocated requests, and zombie cloud instances.",
        "Transforms FinOps from manual sprint tickets into an autonomous background service."
    ], border_color=C_CYAN, title_color=C_CYAN)

    add_card(s2, Inches(4.8), Inches(2.1), Inches(3.6), Inches(4.7), "Predictive & Safe by Design", [
        "Embeds a What-If Simulator that forecasts post-change utilization before touching any cluster.",
        "Strict safety constitutional gatekeeper enforces namespace filters and change windows.",
        "60-second health sentinel with automated rollback if pod restarts or 5xx spikes occur."
    ], border_color=C_BORDER_BLUE, title_color=C_BORDER_BLUE)

    add_card(s2, Inches(8.8), Inches(2.1), Inches(3.7), Inches(4.7), "Collaborative ChatOps & UI", [
        "Slack ChatOps with `/autopilot` commands for team-wide transparency.",
        "Interactive Block Kit cards allow one-click human-in-the-loop sign-off.",
        "3D WebGL cyber cluster map provides spatial visual awareness of infrastructure waste."
    ], border_color=C_GREEN, title_color=C_GREEN)

    # =========================================================================
    # SLIDE 3: PROBLEM STATEMENTS
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_bg(s3)
    add_header(s3, "Problem Statements", "The Multi-Billion Dollar Cloud Waste Epidemic", "Enterprises waste 30-45% of allocated cloud spend due to three systemic operational failures.")

    add_card(s3, Inches(0.8), Inches(2.1), Inches(3.6), Inches(4.7), "1. Set-and-Forget Overprovisioning", [
        "Developers assign generous CPU/RAM requests to avoid latency complaints.",
        "Average Kubernetes cluster runs at only 8-15% actual CPU utilization.",
        "Overprovisioned cores compound as microservice deployments multiply.",
        "Result: Cloud bills skyrocket without delivering reliability value."
    ], border_color=C_ROSE, title_color=C_ROSE)

    add_card(s3, Inches(4.8), Inches(2.1), Inches(3.6), Inches(4.7), "2. Ephemeral & Staging Sprawl", [
        "Test namespaces, PR preview pods, and staging environments are left running 24/7.",
        "Zombie EC2 worker instances run continuously handling zero ingress traffic.",
        "Stateless dev services handling 2 req/minute run 5 to 10 full replicas.",
        "Result: Massive idle compute waste across non-production tiers."
    ], border_color=C_AMBER, title_color=C_AMBER)

    add_card(s3, Inches(8.8), Inches(2.1), Inches(3.7), Inches(4.7), "3. The Remediation Void", [
        "Datadog and CloudWatch show charts but cannot execute fixes.",
        "SREs suffer alert fatigue; rightsizing Jira tickets sit neglected for quarters.",
        "Cron scripts lack context and cause catastrophic outages during traffic spikes.",
        "Result: Endless cost observation with zero automated remediation."
    ], border_color=C_BORDER_BLUE, title_color=C_CYAN)

    # =========================================================================
    # SLIDE 4: SOLUTION
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_bg(s4)
    add_header(s4, "The Solution", "Closed-Loop, Risk-Governed FinOps Automation", "DevOps Autopilot pairs multi-agent reasoning with pre-execution simulation and human safety gates.")

    # Results Banner
    res_card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.0), Inches(11.7), Inches(1.1))
    res_card.fill.solid()
    res_card.fill.fore_color.rgb = C_CARD
    res_card.line.color.rgb = C_GREEN
    res_card.line.width = Pt(1.5)

    tx_res = s4.shapes.add_textbox(Inches(1.0), Inches(2.15), Inches(11.3), Inches(0.8))
    p_rc = tx_res.text_frame.paragraphs[0]
    p_rc.text = "PROVEN IMPACT:   32.8% ➔ 14.2% Cluster Waste Reduction   |   ₹18,450/month Direct Savings   |   0 Incidents"
    p_rc.font.size = Pt(14)
    p_rc.font.bold = True
    p_rc.font.color.rgb = C_GREEN

    # 3 Solution pillars
    add_card(s4, Inches(0.8), Inches(3.4), Inches(3.6), Inches(3.5), "1. Agentic Cognitive Flow", [
        "Specialized agents divide labor: Observe ➔ Reason ➔ Verify ➔ Execute ➔ Report.",
        "Calculates multi-dimensional waste scores across CPU, RAM, and network traffic.",
        "Synthesizes actionable scale-down and instance stop directives."
    ], border_color=C_CYAN, title_color=C_CYAN)

    add_card(s4, Inches(4.8), Inches(3.4), Inches(3.6), Inches(3.5), "2. Pre-Execution What-If", [
        "Simulates post-change utilization before issuing cluster API mutations.",
        "Derives deterministic risk scores (LOW, MEDIUM, HIGH) to safeguard SLOs.",
        "Ensures minimum 35% capacity headroom for burst protection."
    ], border_color=C_BORDER_BLUE, title_color=C_BORDER_BLUE)

    add_card(s4, Inches(8.8), Inches(3.4), Inches(3.7), Inches(3.5), "3. Graduated Autonomy", [
        "70% of low-risk staging actions execute autonomously with zero SRE friction.",
        "30% of sensitive actions routed to Slack for one-click human authorization.",
        "Safe dry-run preview mode protects teams during initial onboarding."
    ], border_color=C_GREEN, title_color=C_GREEN)

    # =========================================================================
    # SLIDE 5: TECHNOLOGY STACK
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_bg(s5)
    add_header(s5, "Technology Stack", "Robust, Modular, and Enterprise-Ready Architecture", "Engineered with Python, modern cloud SDKs, ACID persistence, and rich WebGL visuals.")

    tech_cards = [
        ("Core Backend & API", "Python 3.10+ & Flask", [
            "Flask REST API Gateway with CORS support",
            "Multi-Agent sequential handoff orchestrator",
            "Configurable policy engine (change windows, allowlists)"
        ], C_CYAN),
        ("Cognitive Reasoning", "LLM Function Calling", [
            "OpenAI Tool Calling (GPT-4o / GPT-4o-mini)",
            "Structured schema for scale/stop proposals",
            "Fallback to local deterministic rules (0ms latency)"
        ], C_BORDER_BLUE),
        ("Persistence & Audit", "SQLAlchemy & SQLite", [
            "ACID-compliant relational schema",
            "Resource, OptimizationRun, and Finding models",
            "Referential integrity & indexed audit queries"
        ], C_PURPLE),
        ("Cloud & Infra Adapters", "Kubernetes & AWS", [
            "Kubernetes Python Client (CoreV1, AppsV1, Metrics)",
            "AWS Boto3 SDK for EC2 & CloudWatch",
            "High-fidelity synthetic mock mode for demos"
        ], C_AMBER),
        ("Interfaces & ChatOps", "Three.js & Slack", [
            "Interactive 3D WebGL Cyber Cluster Visualizer",
            "Vanilla CSS & Responsive HTML5 Dashboard",
            "Slack Block Kit API with cryptographic HMAC verification"
        ], C_GREEN)
    ]

    card_w = Inches(2.22)
    spacing = Inches(0.15)
    start_x = Inches(0.8)

    for i, (title, subtitle, bullets, col) in enumerate(tech_cards):
        x = start_x + i * (card_w + spacing)
        add_card(s5, x, Inches(2.1), card_w, Inches(4.7), title, [f"{subtitle}"] + bullets, border_color=col, title_color=col)

    # =========================================================================
    # SLIDE 6: INNOVATION & UNIQUENESS
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_bg(s6)
    add_header(s6, "Innovation / Uniqueness", "What Sets DevOps Autopilot Apart from Traditional FinOps?", "Four core architectural innovations that redefine autonomous infrastructure management.")

    add_card(s6, Inches(0.8), Inches(2.1), Inches(5.6), Inches(2.25), "1. Multi-Agent Role Specialization", [
        "Observer, Optimizer, Safety, Executor, and Reporter operate with strict boundaries.",
        "Separation of concerns prevents conflict of interest: the proposing agent CANNOT execute without the Safety Agent's explicit approval."
    ], border_color=C_CYAN, title_color=C_CYAN)

    add_card(s6, Inches(6.8), Inches(2.1), Inches(5.7), Inches(2.25), "2. The 'What-If' Capacity Simulator", [
        "Mathematically projects post-action utilization BEFORE cluster mutation.",
        "Equation: New Util = Old Util × (Old Replicas / New Replicas).",
        "Categorizes risk into LOW (<65%), MEDIUM (65-80%), and HIGH (>80%)."
    ], border_color=C_BORDER_BLUE, title_color=C_BORDER_BLUE)

    add_card(s6, Inches(0.8), Inches(4.55), Inches(5.6), Inches(2.25), "3. Constitutional Governance Guardrails", [
        "Default namespace allowlist (staging, dev, demo) isolates production workloads.",
        "Production changes strictly locked to 02:00-04:00 AM IST maintenance windows.",
        "Single-run blast radius ceiling: max 50% scale-down and max 10 actions/run."
    ], border_color=C_AMBER, title_color=C_AMBER)

    add_card(s6, Inches(6.8), Inches(4.55), Inches(5.7), Inches(2.25), "4. Dual-Plane Interaction (3D WebGL + Slack)", [
        "Spatial 3D cyber cluster maps workloads as orbital rings with real-time color shifts.",
        "Bidirectional Slack ChatOps empowers SREs to trigger runs and sign off directly in incident channels via interactive buttons."
    ], border_color=C_GREEN, title_color=C_GREEN)

    # =========================================================================
    # SLIDE 7: KEY FEATURES
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_bg(s7)
    add_header(s7, "Key Features", "End-to-End Capabilities of the Platform", "Comprehensive feature suite designed for autonomous operation and SRE confidence.")

    features = [
        ("Multi-Metric Waste Scoring", [
            "Weighted formula evaluating CPU underutilization (60%), RAM waste (30%), and idle network throughput (10%).",
            "Normalized scalar score (0-100) eliminates guesswork."
        ], C_CYAN),
        ("Pre-Execution Simulation", [
            "Calculates post-action headroom before touching deployments.",
            "Rejects any proposal threatening to breach 80% SLO ceiling."
        ], C_BORDER_BLUE),
        ("Graduated Execution Modes", [
            "Default Safe Dry-Run mode previews all actions and INR savings.",
            "Live mode patches K8s Deployments & stops AWS EC2 instances."
        ], C_GREEN),
        ("Interactive 3D Topography", [
            "WebGL 3D visualizer maps cluster pods as glowing color-coded spheres.",
            "Hierarchical Waste Heatmap identifies top cost-draining services."
        ], C_PURPLE),
        ("Slack ChatOps & HITL", [
            "Slash commands `/autopilot run`, `/autopilot status`, `/autopilot explain`.",
            "One-click [Approve & Execute] buttons with HMAC verification."
        ], C_AMBER),
        ("Immutable Audit Ledger", [
            "Records every finding, before/after spec diff, timestamp, and approver.",
            "Full explainability with natural-language agent reasoning traces."
        ], C_ROSE)
    ]

    card_w2 = Inches(3.64)
    card_h2 = Inches(2.25)
    
    positions = [
        (Inches(0.8), Inches(2.1)),
        (Inches(4.8), Inches(2.1)),
        (Inches(8.8), Inches(2.1)),
        (Inches(0.8), Inches(4.55)),
        (Inches(4.8), Inches(4.55)),
        (Inches(8.8), Inches(4.55))
    ]

    for i, (title, bullets, col) in enumerate(features):
        pos_x, pos_y = positions[i]
        add_card(s7, pos_x, pos_y, card_w2, card_h2, title, bullets, border_color=col, title_color=col)

    # =========================================================================
    # SLIDE 8: SUMMARY & DEMO IMPACT (Conclusion)
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_bg(s8)

    bar8 = s8.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.5), Inches(1.5), Inches(0.08))
    bar8.fill.solid()
    bar8.fill.fore_color.rgb = C_GREEN
    bar8.line.fill.background()

    tr8 = s8.shapes.add_textbox(Inches(0.8), Inches(1.75), Inches(11.5), Inches(0.5))
    p_tr8 = tr8.text_frame.paragraphs[0]
    p_tr8.text = "EXECUTIVE SUMMARY  |  HACKATHON RESULTS"
    p_tr8.font.size = Pt(12)
    p_tr8.font.bold = True
    p_tr8.font.color.rgb = C_GREEN

    pn8 = s8.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.7), Inches(1.4))
    tf_pn8 = pn8.text_frame
    tf_pn8.word_wrap = True
    p_pn8 = tf_pn8.paragraphs[0]
    p_pn8.text = "Autonomous FinOps That Engineers Can Trust"
    p_pn8.font.size = Pt(36)
    p_pn8.font.bold = True
    p_pn8.font.color.rgb = C_WHITE

    # 4 Highlight Stats
    stat_data = [
        ("CLUSTER WASTE DROP", "32.8% ➔ 14.2%", "56% Relative Reduction", C_CYAN),
        ("MONTHLY RUN-RATE SAVED", "₹18,450 / month", "₹2,21,400 / year", C_GREEN),
        ("AUTONOMOUS ADOPTION", "70% Auto-Approved", "30% Human Sign-Off", C_BORDER_BLUE),
        ("SLO STABILITY", "0 Outages", "100% Availability Maintained", C_AMBER)
    ]
    sw = Inches(2.78)
    for i, (lbl, val, sub, col) in enumerate(stat_data):
        x = Inches(0.8) + i * (sw + Inches(0.2))
        box = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(3.8), sw, Inches(1.8))
        box.fill.solid()
        box.fill.fore_color.rgb = C_CARD
        box.line.color.rgb = col
        box.line.width = Pt(1.5)

        tb = s8.shapes.add_textbox(x + Inches(0.1), Inches(3.95), sw - Inches(0.2), Inches(1.5))
        tf_b = tb.text_frame
        p1 = tf_b.paragraphs[0]
        p1.text = lbl
        p1.font.size = Pt(10)
        p1.font.bold = True
        p1.font.color.rgb = C_TEXT_MUTED

        p2 = tf_b.add_paragraph()
        p2.text = val
        p2.font.size = Pt(20)
        p2.font.bold = True
        p2.font.color.rgb = col

        p3 = tf_b.add_paragraph()
        p3.text = sub
        p3.font.size = Pt(11)
        p3.font.color.rgb = C_WHITE

    # Links
    links_card = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.9), Inches(11.7), Inches(0.8))
    links_card.fill.solid()
    links_card.fill.fore_color.rgb = C_CARD
    links_card.line.color.rgb = C_BORDER_BLUE
    links_card.line.width = Pt(1)

    tx_l = s8.shapes.add_textbox(Inches(1.0), Inches(6.05), Inches(11.3), Inches(0.5))
    p_l = tx_l.text_frame.paragraphs[0]
    p_l.text = "🚀 Web Dashboard: http://127.0.0.1:5000  |  PDF Report: /report/pdf  |  PPT Pitch Deck: /report/ppt"
    p_l.font.size = Pt(11)
    p_l.font.bold = True
    p_l.font.color.rgb = C_CYAN

    # Save to file
    out_file = os.path.join(os.path.dirname(__file__), "DevOps_Autopilot_Hackathon_Presentation.pptx")
    prs.save(out_file)
    size_kb = os.path.getsize(out_file) / 1024
    print(f"SUCCESS: Generated tailored PPT at: {out_file} ({size_kb:.1f} KB)")
    return out_file

if __name__ == "__main__":
    build_presentation()
